"""Read-only-to-production audit probes. All writes use a temporary DATA_DIR.

Run from backend:
  .venv/bin/python ../docs/evidence/整体链路自查/reproduce.py
No real model calls, search requests, Feishu messages, or live task mutations.
"""
from __future__ import annotations

import asyncio
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import uuid
from unittest.mock import patch

PROJECT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(PROJECT / "backend"))


async def audit(root: Path) -> dict:
    os.environ.update(
        DATA_DIR=str(root / "data"), LOCAL_WORKSPACE_ROOT=str(root / "workspace"),
        LLM_MOCK="1", SEARCH_MOCK="1", FEISHU_MOCK="1",
        COST_SOFT_LIMIT_CNY="5", LLM_PRICE_PER_1K_CNY="0.02",
        FEISHU_NOTIFY_ON="off",
    )
    (root / "workspace").mkdir()
    from app.core.config import get_settings
    get_settings.cache_clear()
    settings = get_settings()
    from app.db import session as ds
    from app.db.models import Task, TaskStep, ScheduleRun, make_engine, make_session_factory
    from app.services import tasks, workspace, schedules, files
    from app.api import tasks as task_api
    from app.core.errors import AppError
    from sqlalchemy import select
    from fastapi import UploadFile
    from docx import Document
    ds.engine = make_engine(f"sqlite:///{settings.db_path}")
    ds.SessionLocal = make_session_factory(ds.engine)
    ds.init_db()
    result = {}

    def make_task(db, status="plan_ready", hints=("read_uploads", "analyze", "write_report")):
        task = Task(
            id=str(uuid.uuid4()), title="AUDIT SYNTHETIC", user_prompt="根据我提供的要点起草简短周报",
            status=status, plan_json=json.dumps({"urls": [], "skill_id": "weekly_report"}),
        )
        db.add(task)
        db.flush()
        for n, hint in enumerate(hints, 1):
            db.add(TaskStep(task_id=task.id, seq=n, name=f"步骤 {n}", status="pending",
                            detail_json=json.dumps({"tool_hint": hint, "goal": "审查"})))
        db.commit()
        return tasks.get_task(db, task.id)

    with ds.SessionLocal() as db:
        marker = "AUDIT_UNSELECTED_WORKSPACE_MARKER"
        (root / "workspace" / "not-selected.txt").write_text(marker, encoding="utf-8")
        task = make_task(db, hints=("read_workspace", "analyze", "write_report"))
        captured = []
        original = tasks.chat_completion
        async def capture(s, system, user):
            captured.append(user)
            return await original(s, system, user)
        with patch.object(tasks, "chat_completion", capture):
            await tasks.confirm_and_run(db, task.id, settings=settings)
        result["unselected_workspace"] = {
            "attached_files": len(task.uploads), "status": tasks.get_task(db, task.id).status,
            "unselected_marker_sent_to_model": any(marker in text for text in captured),
        }

        long_task = make_task(db)
        tail = "AUDIT_TAIL_REQUIRED_FACT_987"
        content = ("前文" * 7000 + tail).encode("utf-8")
        row = await tasks.add_upload(db, long_task.id, UploadFile(filename="long.txt", file=io.BytesIO(content)), settings)
        result["long_text"] = {
            "input_chars": len(content.decode()), "stored_excerpt_chars": len(row.text_excerpt),
            "tail_preserved": tail in row.text_excerpt,
        }

        doc = Document()
        doc.add_paragraph("会议材料")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "AUDIT_DOCX_TABLE_REQUIRED_FACT_654"
        out = io.BytesIO()
        doc.save(out)
        doc_task = make_task(db)
        doc_row = await tasks.add_upload(db, doc_task.id, UploadFile(filename="table.docx", file=io.BytesIO(out.getvalue())), settings)
        result["docx_table"] = {"table_cell_preserved": "AUDIT_DOCX_TABLE_REQUIRED_FACT_654" in doc_row.text_excerpt}

        first = workspace.upload_table_file(filename="same.csv", data=b"x\n1\n", settings=settings)
        second = workspace.upload_table_file(filename="same.csv", data=b"x\n2\n", settings=settings)
        third = workspace.upload_table_file(filename="same.csv", data=b"x\n3\n", settings=settings)
        result["workspace_collision"] = {
            "returned_paths": [first["rel"], second["rel"], third["rel"]],
            "second_upload_overwritten": (root / "workspace" / second["rel"]).read_bytes() != b"x\n2\n",
        }

        planning = make_task(db, status="planning")
        running = make_task(db, status="running")
        recovered = tasks.recover_running_tasks(db)
        db.expire_all()
        result["restart_recovery"] = {
            "recovered_count": recovered, "planning_status": tasks.get_task(db, planning.id).status,
            "running_status": tasks.get_task(db, running.id).status,
        }
        try:
            await tasks.replan_task(db, planning.id, settings=settings)
            result["restart_recovery"]["planning_replan"] = "allowed"
        except AppError as e:
            result["restart_recovery"]["planning_replan"] = e.code

        running2 = make_task(db, status="running")
        replanned = await task_api.rebuild_plan(running2.id, db)
        result["legacy_plan_endpoint"] = {"before": "running", "after": replanned.status}

        schedule = schedules.create_schedule(db, name="AUDIT PAUSE", prompt="写周报", enabled=False, model_id="mock")
        async def pause_instead_of_finish(run_db, task_id, **kwargs):
            task = tasks.get_task(run_db, task_id)
            task.status = "paused"
            run_db.commit()
            return task
        with patch.object(tasks, "confirm_and_run", pause_instead_of_finish):
            await schedules.trigger_schedule(db, schedule.id)
        run = schedules.list_schedule_runs(db, schedule.id)[0]
        result["paused_schedule"] = {
            "task_status": tasks.get_task(db, run.task_id).status,
            "run_status": run.status, "run_error": run.error,
        }

        schedule2 = schedules.create_schedule(db, name="AUDIT CONCURRENT", prompt="写周报", enabled=False, model_id="mock")
        original_create = tasks.create_task
        async def yield_create(*args, **kwargs):
            await asyncio.sleep(0.04)
            return await original_create(*args, **kwargs)
        async def invoke():
            with ds.SessionLocal() as other:
                await schedules.trigger_schedule(other, schedule2.id)
        with patch.object(tasks, "create_task", yield_create):
            await asyncio.gather(invoke(), invoke())
        db.expire_all()
        runs = schedules.list_schedule_runs(db, schedule2.id)
        result["concurrent_schedule"] = {
            "runs_created": len(runs), "distinct_tasks": len({r.task_id for r in runs}),
            "statuses": [r.status for r in runs],
        }
    ds.engine.dispose()
    get_settings.cache_clear()
    return result


if __name__ == "__main__":
    with tempfile.TemporaryDirectory(prefix="office-chain-audit-") as folder:
        print(json.dumps(asyncio.run(audit(Path(folder))), ensure_ascii=False, indent=2))
