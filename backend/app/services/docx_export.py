from __future__ import annotations

import re
from pathlib import Path

from docx import Document


def markdown_to_docx(md: str, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        else:
            doc.add_paragraph(line)
    # python-docx has no atomic helper; write via temp then replace
    tmp = dest.with_suffix(dest.suffix + ".tmp")
    doc.save(tmp)
    tmp.replace(dest)
    return dest


def ensure_docx_from_markdown(md_path: Path, docx_path: Path) -> Path:
    md = md_path.read_text(encoding="utf-8")
    return markdown_to_docx(md, docx_path)
