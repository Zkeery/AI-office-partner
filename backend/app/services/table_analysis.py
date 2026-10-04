"""Compute on all accepted rows. Model prompts receive statistics, not a sampled table."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import statistics
import uuid
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.errors import AppError
from app.services.files import atomic_write_text

MAX_ROWS = 100_000
MAX_CELLS = 1_000_000
MAX_COLUMNS = 512
MAX_SHEETS = 50
ID_PATTERN = re.compile(r"(^id$|编号|编码|邮编|手机|电话|身份证|序号)", re.I)
METRIC_PATTERN = re.compile(r"金额|销售额|收入|营收|利润|成本|数量|销量|单价|费用|支出|率|amount|revenue|sales|price|quantity|cost", re.I)


def _cell(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    return str(value)


def number(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1].strip()
    text = re.sub(r"^[¥￥$]\s*", "", text)
    percent = text.endswith("%") or text.endswith("％")
    text = text.rstrip("%％").strip()
    if "," in text and not re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", text):
        return None
    text = text.replace(",", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        return None
    try:
        parsed = Decimal(text)
        if not parsed.is_finite() or abs(parsed) > Decimal("1e100"):
            return None
        if negative:
            parsed = -parsed
        return parsed / 100 if percent else parsed
    except InvalidOperation:
        return None


def _round(value: Decimal | float | None) -> float | None:
    return round(float(value), 10) if value is not None else None


def month(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    match = re.match(r"^(\d{4})[-/年](\d{1,2})(?:[-/月](\d{1,2})日?)?(?:$|[T ])", value.strip())
    if not match:
        return None
    try:
        parsed = date(int(match[1]), int(match[2]), int(match[3] or 1))
        return f"{parsed.year:04d}-{parsed.month:02d}"
    except ValueError:
        return None


def _empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _headers(values: list[Any]) -> list[str]:
    result: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if not _empty(value) else f"列{index}"
        candidate, suffix = base, 2
        while candidate in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        result.append(candidate)
        used.add(candidate)
    return result


def read_tables(path: Path, source: str, budget: dict[str, int]) -> list[dict[str, Any]]:
    """The first nonempty row is the header. Never silently truncate accepted data."""
    tables: list[dict[str, Any]] = []

    def consume(rows, sheet: str) -> None:
        if budget["sheets"] >= MAX_SHEETS:
            raise AppError("TABLE_TOO_LARGE", "一次最多分析 50 个工作表，请拆分文件后重试")
        raw_header: list[Any] | None = None
        records: list[list[Any]] = []
        row_numbers: list[int] = []
        width = 0
        for row_number, raw in rows:
            values = list(raw)
            while values and _empty(values[-1]):
                values.pop()
            if not values:
                continue
            if len(values) > MAX_COLUMNS:
                raise AppError("TABLE_TOO_LARGE", f"「{source} / {sheet}」超过 512 列，请拆分文件")
            for value in values:
                if isinstance(value, str) and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
                    raise AppError("TABLE_CELL_INVALID", f"「{source} / {sheet}」第 {row_number} 行含过长文本或非法控制字符，无法完整导出 Excel，请清理后上传")
            budget["cells"] += len(values)
            if budget["cells"] > MAX_CELLS:
                raise AppError("TABLE_TOO_LARGE", "本次表格超过 100 万个单元格，请拆分后分析；未使用抽样结果")
            width = max(width, len(values))
            if raw_header is None:
                raw_header = values
                continue
            budget["rows"] += 1
            if budget["rows"] > MAX_ROWS:
                raise AppError("TABLE_TOO_LARGE", "本次表格超过 10 万行，请拆分后分析；未使用抽样结果")
            records.append([_cell(value) for value in values])
            row_numbers.append(row_number)
        if raw_header is None:
            return
        budget["sheets"] += 1
        headers = _headers(raw_header + [None] * (width - len(raw_header)))
        tables.append({
            "id": f"table-{budget['sheets']}",
            "source": source,
            "sheet": sheet,
            "headers": headers,
            "rows": [row + [None] * (width - len(row)) for row in records],
            "row_numbers": row_numbers,
        })

    if path.suffix.lower() == ".csv":
        data = path.read_bytes()
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise AppError("TABLE_ENCODING", "CSV 编码无法识别，请另存为 UTF-8 CSV") from exc
        if "\x00" in text:
            raise AppError("TABLE_INVALID", "CSV 含有二进制内容，请上传有效表格")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        try:
            consume(enumerate(csv.reader(io.StringIO(text), dialect), start=1), path.stem if source == path.name else Path(source).stem)
        except csv.Error as exc:
            raise AppError("TABLE_INVALID", "CSV 格式不完整，请检查引号和分隔符") from exc
        return tables

    if path.suffix.lower() != ".xlsx":
        raise AppError("TABLE_TYPE", "数据分析支持 CSV 和 XLSX 文件")
    values_book = formula_book = None
    try:
        values_book = load_workbook(path, read_only=True, data_only=True)
        formula_book = load_workbook(path, read_only=True, data_only=False)
        for sheet in values_book.worksheets:
            formulas = formula_book[sheet.title]
            if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 1:
                raise AppError("TABLE_TOO_LARGE", f"「{sheet.title}」超出单表容量（10 万行、512 列），请清理空白格式区域或拆分表格")

            def checked_rows():
                for row_number, (cached, raw) in enumerate(zip(sheet.iter_rows(), formulas.iter_rows()), start=1):
                    for index, cell in enumerate(raw):
                        if cell.data_type == "f" and (index >= len(cached) or cached[index].value is None):
                            raise AppError("TABLE_FORMULA_UNCALCULATED", f"「{sheet.title}!{cell.coordinate}」的公式没有计算结果。请在 Excel 中重新计算并保存，或粘贴为数值后上传")
                    yield row_number, [cell.value for cell in cached]

            consume(checked_rows(), sheet.title)
        return tables
    except AppError:
        raise
    except Exception as exc:
        raise AppError("TABLE_INVALID", f"「{source}」无法作为 XLSX 读取，请确认文件未损坏或加密") from exc
    finally:
        if values_book:
            values_book.close()
        if formula_book:
            formula_book.close()


def _preferred_metric(metrics: list[dict[str, Any]], prompt: str) -> dict[str, Any] | None:
    if not metrics:
        return None
    named = [m for m in metrics if m["column"] in prompt]
    if named:
        return named[0]
    priority = re.compile(r"销售额|收入|营收|金额|revenue|amount|sales", re.I)
    return next((m for m in metrics if priority.search(m["column"])), metrics[0])


def summarize(table: dict[str, Any], prompt: str) -> dict[str, Any]:
    rows, headers = table["rows"], table["headers"]
    metrics: list[dict[str, Any]] = []
    numeric: dict[int, list[Decimal | None]] = {}
    issues: list[dict[str, Any]] = []
    issue_count = 0
    warnings: list[str] = []
    date_index: int | None = None
    categorical: list[int] = []

    def add_issue(index: int, row: int, kind: str, message: str) -> None:
        nonlocal issue_count
        issue_count += 1
        if len(issues) < 100:
            issues.append({"row": table["row_numbers"][row], "column": headers[index], "kind": kind, "value": str(rows[row][index] if rows[row][index] is not None else "")[:200], "message": message})

    for index, header in enumerate(headers):
        values = [row[index] for row in rows]
        nonempty = [value for value in values if not _empty(value)]
        dates = sum(month(value) is not None for value in nonempty)
        if nonempty and dates / len(nonempty) >= 0.8:
            if date_index is None:
                date_index = index
            continue
        parsed = [number(value) for value in values]
        valid = [value for value in parsed if value is not None]
        id_like = bool(ID_PATTERN.search(header)) or any(isinstance(value, str) and re.fullmatch(r"0\d+", value.strip()) for value in nonempty)
        numeric_field = bool(valid) and not id_like and (len(valid) / max(len(nonempty), 1) >= 0.8 or bool(METRIC_PATTERN.search(header)))
        if not numeric_field:
            labels = {str(value) for value in nonempty}
            if labels and len(labels) <= 100 and not id_like:
                categorical.append(index)
            continue
        numeric[index] = parsed
        invalid = sum(not _empty(value) and parsed[i] is None for i, value in enumerate(values))
        missing = len(values) - len(nonempty)
        metric = {
            "column": header, "index": index, "kind": "percent" if any(isinstance(v, str) and v.strip().endswith(("%", "％")) for v in nonempty) else "number",
            "count": len(valid), "missing_count": missing, "invalid_count": invalid,
            "sum": _round(sum(valid, Decimal(0))), "mean": _round(sum(valid, Decimal(0)) / len(valid)),
            "min": _round(min(valid)), "max": _round(max(valid)), "outlier_count": 0,
        }
        metric["aggregation"] = "mean" if metric["kind"] == "percent" or re.search(r"单价|均价|平均|率$|ratio|rate|price|average", header, re.I) else "sum"
        if invalid or missing:
            warnings.append(f"「{header}」有 {missing} 个空值、{invalid} 个非数字值，统计仅使用 {len(valid)} 个有效数字")
        bounds = None
        if len(valid) >= 4:
            q1, _, q3 = statistics.quantiles(valid, n=4, method="inclusive")
            spread = q3 - q1
            bounds = (q1 - Decimal("1.5") * spread, q3 + Decimal("1.5") * spread)
        for row, value in enumerate(values):
            if not _empty(value) and parsed[row] is None:
                add_issue(index, row, "invalid", "不能作为数字计算，已排除")
            elif bounds and parsed[row] is not None and (parsed[row] < bounds[0] or parsed[row] > bounds[1]):
                metric["outlier_count"] += 1
                add_issue(index, row, "outlier", "超出四分位距的 1.5 倍范围，请核对；仍计入统计")
        metrics.append(metric)

    preferred = _preferred_metric(metrics, prompt)
    groups: list[dict[str, Any]] = []
    charts: list[dict[str, Any]] = []
    if preferred:
        metric_index = preferred["index"]
        for dim in categorical[:3]:
            sums: dict[str, Decimal] = defaultdict(Decimal)
            counts: dict[str, int] = defaultdict(int)
            for row_index, row in enumerate(rows):
                value = numeric[metric_index][row_index]
                if value is None:
                    continue
                label = "（空白）" if _empty(row[dim]) else str(row[dim])
                sums[label] += value
                counts[label] += 1
            grouped = [{"label": label, "value": _round(total / counts[label] if preferred["aggregation"] == "mean" else total), "count": counts[label]} for label, total in sums.items()]
            grouped.sort(key=lambda row: (-row["value"], row["label"]))
            groups.append({"dimension": headers[dim], "metric": preferred["column"], "aggregation": preferred["aggregation"], "rows": grouped})
        if groups:
            group = groups[0]
            shown = group["rows"][:12]
            label = "均值（未加权）" if group["aggregation"] == "mean" else "汇总"
            charts.append({"kind": "bar", "title": f"{group['metric']} · 按{group['dimension']}{label}" + ("（前 12 项）" if len(group["rows"]) > 12 else ""), "labels": [r["label"] for r in shown], "values": [r["value"] for r in shown], "metric": group["metric"]})

    periods: list[dict[str, Any]] = []
    if date_index is not None:
        for metric in metrics:
            totals: dict[str, Decimal] = defaultdict(Decimal)
            counts: dict[str, int] = defaultdict(int)
            invalid_dates = 0
            for row_index, row in enumerate(rows):
                value = numeric[metric["index"]][row_index]
                if value is None:
                    continue
                period = month(row[date_index])
                if period is None:
                    invalid_dates += 1
                    continue
                totals[period] += value
                counts[period] += 1
            if invalid_dates:
                warnings.append(f"「{metric['column']}」有 {invalid_dates} 行日期缺失或无法识别，未计入月度汇总")
            if metric["aggregation"] == "mean":
                totals = {period: total / counts[period] for period, total in totals.items()}
            for period, total in sorted(totals.items()):
                year, mon = (int(part) for part in period.split("-"))
                prev_key = f"{year - 1}-12" if mon == 1 else f"{year}-{mon - 1:02d}"
                previous, last_year = totals.get(prev_key), totals.get(f"{year - 1}-{mon:02d}")
                periods.append({
                    "period": period, "metric": metric["column"], "aggregation": metric["aggregation"], "value": _round(total),
                    "mom_pct": _round((total - previous) / abs(previous) * 100) if previous else None,
                    "yoy_pct": _round((total - last_year) / abs(last_year) * 100) if last_year else None,
                })
        if preferred:
            series = [period for period in periods if period["metric"] == preferred["column"]]
            if series:
                charts.append({"kind": "line", "title": f"{preferred['column']} · 月度趋势" + ("（均值，未加权）" if preferred["aggregation"] == "mean" else ""), "labels": [r["period"] for r in series], "values": [r["value"] for r in series], "metric": preferred["column"]})

    if not rows:
        warnings.append("此工作表只有表头，没有数据行")
    if rows and not metrics:
        warnings.append("没有识别到可计算的数字列；保留完整数据，不生成数字结论")
    for chart in charts:
        chart["value_format"] = preferred["kind"] if preferred else "number"
    return {
        "id": table["id"], "source": table["source"], "sheet": table["sheet"],
        "row_count": len(rows), "column_count": len(headers), "headers": headers,
        "metrics": metrics, "groups": groups, "periods": periods,
        "date_column": headers[date_index] if date_index is not None else None,
        "issues": issues, "issue_count": issue_count, "warnings": warnings, "charts": charts,
    }


def build_snapshot(uploads: list[dict[str, str]], artifact_dir: Path, prompt: str) -> str:
    budget = {"rows": 0, "cells": 0, "sheets": 0}
    tables: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    for upload in uploads:
        path = Path(upload["stored_path"])
        if not path.is_file():
            raise AppError("TABLE_FILE_MISSING", f"找不到材料「{upload['filename']}」，请重新上传")
        sources.append({"filename": upload["filename"], "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
        tables.extend(read_tables(path, upload["filename"], budget))
    if not tables:
        raise AppError("TABLE_EMPTY", "表格没有可读取的表头或数据，请检查文件")
    analysis_id = uuid.uuid4().hex
    summary = {
        "id": analysis_id, "complete": True, "created_at": datetime.now(timezone.utc).isoformat(),
        "source_count": len(sources), "sheet_count": len(tables), "row_count": budget["rows"],
        "cell_count": budget["cells"], "sources": sources,
        "tables": [summarize(table, prompt) for table in tables],
    }
    atomic_write_text(artifact_dir / "tables" / f"{analysis_id}.json", json.dumps({"summary": summary, "datasets": tables}, ensure_ascii=False, allow_nan=False))
    return analysis_id


def load_snapshot(artifact_dir: Path, analysis_id: str) -> dict[str, Any]:
    if not re.fullmatch(r"[a-f0-9]{32}", analysis_id):
        raise AppError("ANALYSIS_NOT_FOUND", "数据分析记录不存在", status_code=404)
    path = artifact_dir / "tables" / f"{analysis_id}.json"
    if not path.is_file():
        raise AppError("ANALYSIS_NOT_FOUND", "数据分析记录不存在，请重新运行任务", status_code=404)
    return json.loads(path.read_text(encoding="utf-8"))


def prompt_summary(summary: dict[str, Any]) -> str:
    """Bound prompt size explicitly while retaining all results in the local workbook."""
    tables = []
    for table in summary["tables"]:
        tables.append({
            "文件": table["source"], "工作表": table["sheet"], "完整数据行数": table["row_count"],
            "总列数": table["column_count"], "程序计算指标": table["metrics"][:20],
            "分组结果": [{**group, "rows": group["rows"][:20]} for group in table["groups"]],
            "月度汇总": table["periods"][:120], "注意": table["warnings"],
            "实际生成图表": [{"标题": chart["title"], "指标": chart["metric"]} for chart in table["charts"]],
            "异常总数": table["issue_count"], "异常示例": table["issues"][:5],
        })
    content = json.dumps({"完整读取": True, "全部行数": summary["row_count"], "全部工作表数": summary["sheet_count"], "统计": tables}, ensure_ascii=False)
    if len(content) > 45000:
        content = json.dumps({"完整读取": True, "全部行数": summary["row_count"], "全部工作表数": summary["sheet_count"], "统计": [{"文件": t["source"], "工作表": t["sheet"], "行数": t["row_count"], "指标": t["metrics"][:5], "实际生成图表": [c["title"] for c in t["charts"]], "注意": t["warnings"][:5]} for t in summary["tables"]]}, ensure_ascii=False)
    return (
        f"本次实际生成 {sum(len(t['charts']) for t in summary['tables'])} 张图表，须逐工作表核对清单，不能遗漏其他工作表的图表。"
        "以下结果由程序对全部已接收数据计算。只解释已提供的指标；未提供的计算不得猜测。"
        "空值和非数字排除，离群值仍计入，环比同比缺失或基期为 0 时为 null。"
        "汇总求和不等于业务口径已确认，比例、单价等指标优先说明未加权均值；整体转化率需原始分子和分母，不能猜测权重。"
        "只有‘实际生成图表’清单中的图才已生成，其他统计不得称为已有图表。"
        "正文将 null 写成‘—（缺少可用基期）’，将百分比指标的小数转换为百分数。"
        "正文使用中文指标名与自然语言，不暴露字段名、analysis_id、哈希等实现细节；来源写文件名与工作表名。"
        "完整明细、公式、图表见 Excel；以下仅压缩展示已算出的结果：\n"
    ) + content
