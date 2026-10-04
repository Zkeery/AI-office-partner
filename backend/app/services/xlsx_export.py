from __future__ import annotations

import re
import math
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.properties import CalcProperties

from app.services.table_analysis import number

INK = "173B45"
TEAL = "167D8D"
PALE = "EAF4F4"


def _append(ws, values: list[Any], *, formulas: bool = False) -> None:
    ws.append(values)
    if not formulas:
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"


def _header(ws, values: list[str]) -> None:
    _append(ws, values)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def _finish_sheet(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _save(wb: Workbook, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    tmp = dest.with_name(f".{dest.stem}-{uuid.uuid4().hex}.xlsx")
    try:
        wb.save(tmp)
        tmp.replace(dest)
    finally:
        tmp.unlink(missing_ok=True)
        wb.close()
    return dest


def _report_sheet(wb: Workbook, md: str) -> None:
    ws = wb.active
    ws.title = "报告"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110
    for raw in md.splitlines():
        kind, content = _row_kind(raw.rstrip())
        if not content.strip():
            continue
        for start in range(0, len(content), 1200):
            piece = content[start:start + 1200]
            _append(ws, [piece])
            cell = ws.cell(ws.max_row, 1)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Aptos", size=18 if kind == "标题1" else 13 if kind.startswith("标题") else 11, bold=kind.startswith("标题"), color=INK)
            if kind.startswith("标题"):
                cell.fill = PatternFill("solid", fgColor=PALE)
            ws.row_dimensions[ws.max_row].height = 20 + math.ceil(len(piece) / 75) * 16
    ws.freeze_panes = "A2"


def _native_chart(ws, *, kind: str, title: str, start: int, end: int, label_col: int, value_col: int, anchor: str, value_format: str = "number") -> None:
    if end < start:
        return
    chart = LineChart() if kind == "line" else BarChart()
    chart.title = title
    chart.style = 13 if kind == "line" else 10
    chart.add_data(Reference(ws, min_col=value_col, min_row=start, max_row=end))
    chart.set_categories(Reference(ws, min_col=label_col, min_row=start, max_row=end))
    chart.legend = None
    chart.height, chart.width = 9, 18
    chart.y_axis.title = "数值"
    values = [ws.cell(row, value_col).value for row in range(start, end + 1)]
    chart.y_axis.scaling.min = min(0, *values)
    if max(values) <= 0:
        chart.y_axis.scaling.max = 0 if min(values) < 0 else 1
    if value_format == "percent":
        chart.y_axis.numFmt = "0.0%"
    ws.add_chart(chart, anchor)


def _row_kind(line: str) -> tuple[str, str]:
    if line.startswith("### "):
        return "标题3", line[4:].strip()
    if line.startswith("## "):
        return "标题2", line[3:].strip()
    if line.startswith("# "):
        return "标题1", line[2:].strip()
    if re.match(r"^[-*]\s+", line):
        return "要点", re.sub(r"^[-*]\s+", "", line).strip()
    if re.match(r"^\d+\.\s+", line):
        return "要点", re.sub(r"^\d+\.\s+", "", line).strip()
    return "正文", line


def markdown_to_xlsx(md: str, dest: Path) -> Path:
    wb = Workbook()
    _report_sheet(wb, md)
    # Markdown tables also become editable cells instead of a single text cell.
    blocks = re.findall(r"(?m)^\|.+\|\s*\n\|[\s:|\-]+\|\s*\n(?:\|.+\|[ \t]*\n?)+", md)
    for index, block in enumerate(blocks, start=1):
        lines = [line.strip().strip("|").split("|") for line in block.strip().splitlines()]
        ws = wb.create_sheet(f"表格{index}")
        _header(ws, [value.strip() for value in lines[0]])
        for line in lines[2:]:
            values = [value.strip() for value in line]
            _append(ws, [float(number(value)) if number(value) is not None else value for value in values])
        _finish_sheet(ws, [24] * len(lines[0]))
    return _save(wb, dest)


def analysis_to_xlsx(snapshot: dict[str, Any], md: str, dest: Path) -> Path:
    wb = Workbook()
    _report_sheet(wb, md)
    summary = snapshot["summary"]
    overview = wb.create_sheet("数据口径")
    _header(overview, ["项目", "说明"])
    for row in [
        ["读取范围", f"{summary['source_count']} 个文件，{summary['sheet_count']} 个工作表，{summary['row_count']} 行；所有已接收数据均参与计算"],
        ["表头", "每张表第一条非空行作为表头；全空行不计入数据行"],
        ["计算口径", "空值和非数字不计入数值统计；离群值保留在统计中"],
        ["环比 / 同比", "与上一个自然月 / 上一年同月比较；基期缺失或为 0 时留空"],
        ["复算公式", "指标汇总含 SUM 或 AVERAGE 公式；旁边为程序计算的快照值，修改原始数据后需重新上传分析以更新图表与结论"],
        ["异常明细", "每个工作表最多展示前 100 条异常，异常总数见下方；原始数据完整保留"],
        ["生成时间", summary["created_at"]],
    ]:
        _append(overview, row)
    _finish_sheet(overview, [24, 110])

    stats = wb.create_sheet("指标汇总")
    _header(stats, ["文件", "工作表", "指标", "有效数", "空值", "非数字", "求和", "均值", "最小值", "最大值", "Excel复算", "分组/月度口径"])
    grouped = wb.create_sheet("分组统计")
    _header(grouped, ["文件", "工作表", "分组字段", "分类", "指标", "计算值", "有效数", "计算口径"])
    monthly = wb.create_sheet("月度趋势")
    _header(monthly, ["文件", "工作表", "月份", "指标", "计算值", "环比", "同比", "计算口径"])
    anomalies = wb.create_sheet("异常明细")
    _header(anomalies, ["文件", "工作表", "原始行号", "字段", "值", "类型", "说明"])
    group_charts = month_charts = 0

    for table, data in zip(summary["tables"], snapshot["datasets"]):
        base = re.sub(r"[\[\]:*?/\\]", "_", f"数据-{table['sheet']}")[:31]
        title, suffix = base, 2
        while title in wb.sheetnames:
            title = f"{base[:27]}-{suffix}"
            suffix += 1
        raw = wb.create_sheet(title)
        _header(raw, data["headers"])
        metrics = {metric["index"]: metric for metric in table["metrics"]}
        for row in data["rows"]:
            values = list(row)
            for index in metrics:
                parsed = number(values[index])
                if parsed is not None:
                    values[index] = float(parsed)
            _append(raw, values)
        _finish_sheet(raw, [min(38, max(16, len(header) * 2 + 6)) for header in data["headers"]])
        for index, metric in metrics.items():
            col = get_column_letter(index + 1)
            for cells in raw.iter_cols(min_col=index + 1, max_col=index + 1, min_row=2):
                for cell in cells:
                    cell.number_format = "0.00%" if metric["kind"] == "percent" else "#,##0.00;[Red](#,##0.00);–"
            average = metric.get("aggregation") == "mean"
            formula = f"={'AVERAGE' if average else 'SUM'}({quote_sheetname(title)}!{col}2:{col}{max(raw.max_row, 2)})"
            _append(stats, [table["source"], table["sheet"], metric["column"], metric["count"], metric["missing_count"], metric["invalid_count"], None if average else metric["sum"], metric["mean"], metric["min"], metric["max"], None, "均值（未加权）" if average else "求和"])
            stats.cell(stats.max_row, 11, formula)
            for column in range(7, 12):
                stats.cell(stats.max_row, column).number_format = "0.00%" if metric["kind"] == "percent" else "#,##0.00;[Red](#,##0.00);–"
        for index, group in enumerate(table["groups"]):
            start = grouped.max_row + 1
            for row in group["rows"]:
                _append(grouped, [table["source"], table["sheet"], group["dimension"], row["label"], group["metric"], row["value"], row["count"], "均值（未加权）" if group.get("aggregation") == "mean" else "求和"])
            if index == 0 and group["rows"]:
                spec = next(chart for chart in table["charts"] if chart["kind"] == "bar")
                _native_chart(grouped, kind="bar", title=f"{table['sheet']} · {spec['title']}", start=start, end=min(start + 11, grouped.max_row), label_col=4, value_col=6, anchor=f"J{2 + group_charts * 19}", value_format=spec.get("value_format", "number"))
                group_charts += 1
        for metric in table["metrics"]:
            periods = [row for row in table["periods"] if row["metric"] == metric["column"]]
            start = monthly.max_row + 1
            for row in periods:
                _append(monthly, [table["source"], table["sheet"], row["period"], row["metric"], row["value"], row["mom_pct"] / 100 if row["mom_pct"] is not None else None, row["yoy_pct"] / 100 if row["yoy_pct"] is not None else None, "均值（未加权）" if row.get("aggregation") == "mean" else "求和"])
                if metric["kind"] == "percent":
                    monthly.cell(monthly.max_row, 5).number_format = "0.00%"
                monthly.cell(monthly.max_row, 6).number_format = "0.00%;[Red](0.00%);–"
                monthly.cell(monthly.max_row, 7).number_format = "0.00%;[Red](0.00%);–"
            if periods and any(chart["kind"] == "line" and chart["metric"] == metric["column"] for chart in table["charts"]):
                _native_chart(monthly, kind="line", title=f"{table['sheet']} · {metric['column']}", start=start, end=monthly.max_row, label_col=3, value_col=5, anchor=f"J{2 + month_charts * 19}", value_format=metric["kind"])
                month_charts += 1
        _append(overview, [f"{table['source']} / {table['sheet']}", f"{table['row_count']} 行 × {table['column_count']} 列；异常 {table['issue_count']} 条"])
        for warning in table["warnings"]:
            _append(overview, ["数据提示", warning])
        for issue in table["issues"]:
            _append(anomalies, [table["source"], table["sheet"], issue["row"], issue["column"], issue["value"], issue["kind"], issue["message"]])
    for ws in (stats, grouped, monthly, anomalies):
        _finish_sheet(ws, [24] * ws.max_column)
    return _save(wb, dest)


def ensure_xlsx_from_markdown(md_path: Path, xlsx_path: Path) -> Path:
    md = md_path.read_text(encoding="utf-8")
    return markdown_to_xlsx(md, xlsx_path)
