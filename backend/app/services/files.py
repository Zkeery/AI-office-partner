from __future__ import annotations

import csv
import io
import re
from itertools import islice
from pathlib import Path
from typing import Any

from fastapi import UploadFile

from app.core.errors import AppError

ALLOWED_EXT = {".md", ".txt", ".markdown", ".pdf", ".docx", ".csv", ".xlsx"}
MAX_BYTES = 20 * 1024 * 1024
MAX_FILES = 3
MAX_DOCUMENT_CHARS = 60000
MAX_REFERENCE_CHARS = 90000


def safe_filename(name: str) -> str:
    base = Path(name).name
    base = re.sub(r"[^\w.\u4e00-\u9fff\-]+", "_", base).strip("._")
    return base or "upload.bin"


def assert_allowed(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise AppError("UPLOAD_TYPE_NOT_ALLOWED", f"不支持的文件类型：{ext or '无扩展名'}")
    return ext


async def read_upload_bytes(file: UploadFile, max_bytes: int = MAX_BYTES) -> bytes:
    data = await file.read()
    if len(data) > max_bytes:
        raise AppError("UPLOAD_TOO_LARGE", "单个文件不能超过 20MB")
    return data


def _sheet_to_text(rows: list[list[Any]], max_rows: int = 40, max_cols: int = 12) -> str:
    lines: list[str] = []
    for i, row in enumerate(rows[:max_rows]):
        cells = [str(c if c is not None else "").strip() for c in row[:max_cols]]
        if not any(cells):
            continue
        lines.append(" | ".join(cells))
    if len(rows) > max_rows:
        lines.append(f"…（仅展示前 {max_rows} 行）")
    return "\n".join(lines)


def extract_csv_text(data: bytes, max_chars: int = 12000) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")
    reader = csv.reader(io.StringIO(text))
    rows = list(islice(reader, 6))
    return ("以下仅为材料预览，执行时由程序读取完整表格：\n" + _sheet_to_text(rows))[:max_chars]


def extract_xlsx_text(path: Path, max_chars: int = 12000) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = ["以下仅为材料预览，执行时由程序读取全部工作表与数据行："]
        for ws in wb.worksheets:
            rows = [list(row) for row in islice(ws.iter_rows(values_only=True), 6)]
            parts.append(f"工作表：{ws.title}\n" + _sheet_to_text(rows))
        return "\n\n".join(parts)[:max_chars]
    finally:
        wb.close()


def _complete_document(text: str, limit: int) -> str:
    if not text.strip():
        raise AppError("DOCUMENT_EMPTY", "文件没有可读取正文；扫描件请先转成文字，当前不支持 OCR")
    if len(text) > limit:
        raise AppError("DOCUMENT_TOO_LARGE", f"正文共 {len(text):,} 字符，超过单份材料 {limit:,} 字符上限，请拆分后重试；未截取或提交部分正文")
    return text


def extract_text(path: Path, ext: str, max_chars: int = MAX_DOCUMENT_CHARS) -> str:
    """Documents are complete within an explicit bound; table text is a labelled preview."""
    if ext in {".txt", ".md", ".markdown"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        return _complete_document(text, max_chars)
    if ext == ".csv":
        return extract_csv_text(path.read_bytes(), max_chars=max_chars)
    if ext == ".xlsx":
        return extract_xlsx_text(path, max_chars=max_chars)
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        has_text = False
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            has_text = has_text or bool(text.strip())
            parts.append(text if text.strip() else f"[第 {number} 页无可提取文字，可能为空白页或图片；本页未进行 OCR]")
        if not has_text:
            raise AppError("DOCUMENT_EMPTY", "PDF 没有可读取文字；扫描件请先进行 OCR")
        return _complete_document("\n".join(parts), max_chars)
    if ext == ".docx":
        import docx

        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = docx.Document(str(path))
        def blocks(parent, elements):
            parts = []
            for element in elements:
                if element.tag.endswith("}p"):
                    parts.append(Paragraph(element, parent).text)
                elif element.tag.endswith("}tbl"):
                    table = Table(element, parent)
                    parts.append("\n".join(" | ".join(blocks(cell, cell._tc) for cell in row.cells) for row in table.rows))
            return "\n".join(parts)
        return _complete_document(blocks(doc, doc.element.body), max_chars)
    return ""


def task_document_context(uploads) -> str:
    parts = []
    for upload in uploads:
        ext = Path(upload.filename).suffix.lower()
        if ext in {".csv", ".xlsx"}:
            continue
        path = Path(upload.stored_path)
        if not path.is_file():
            raise AppError("MATERIAL_MISSING", f"材料「{upload.filename}」缺失，请重新添加")
        try:
            text = extract_text(path, ext)
        except AppError:
            raise
        except Exception as exc:
            raise AppError("UPLOAD_INVALID", f"材料「{upload.filename}」无法读取，请检查格式或加密状态") from exc
        parts.append(f"【材料：{upload.filename}】\n{text}")
    result = "\n---\n".join(parts)
    if len(result) > MAX_REFERENCE_CHARS:
        raise AppError("MATERIALS_TOO_LARGE", f"本次材料合计超过 {MAX_REFERENCE_CHARS:,} 字符，请分成多个任务；未截断材料")
    return result


def join_context(parts: list[str], max_chars: int = 120000) -> str:
    """Reuse repeated source blocks once, never discard the beginning/end silently."""
    result = "\n---\n".join(dict.fromkeys(part for part in parts if part.strip()))
    if len(result) > max_chars:
        raise AppError("CONTEXT_TOO_LARGE", "任务上下文超过处理上限，请减少材料或拆分任务；未截断材料")
    return result


def atomic_write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(content, encoding="utf-8")
    tmp.replace(path)
