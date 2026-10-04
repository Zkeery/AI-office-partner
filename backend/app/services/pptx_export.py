from __future__ import annotations

import re
import math
import textwrap
import uuid
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

INK = RGBColor.from_string("173B45")
TEAL = RGBColor.from_string("167D8D")


def _heading(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(11.9), Inches(1.1))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(26 if len(title) < 45 else 20)
    paragraph.font.bold = True
    paragraph.font.color.rgb = INK


def _footer(slide, number: int, source: str = "") -> None:
    box = slide.shapes.add_textbox(Inches(0.75), Inches(7.02), Inches(11.8), Inches(0.3))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = (source[:100] + ("…" if len(source) > 100 else "") + "   ·   " if source else "") + f"AI办公搭子  /  {number:02d}"
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = RGBColor.from_string("637A80")


def _flush_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    pieces = [piece for item in (bullets or ["（本节无摘要条目）"]) for piece in textwrap.wrap(item, width=160, replace_whitespace=False, drop_whitespace=False)]
    pages: list[list[str]] = []
    current: list[str] = []
    height = 0
    for piece in pieces:
        lines = max(1, math.ceil(len(piece) / 42)) + 1
        if current and height + lines > 13:
            pages.append(current)
            current, height = [], 0
        current.append(piece)
        height += lines
    if current:
        pages.append(current)
    for index, items in enumerate(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _heading(slide, (title or "内容") + (f"（续 {index + 1}）" if index else ""))
        body = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.65), Inches(4.9)).text_frame
        body.word_wrap = True
        for item_index, item in enumerate(items):
            paragraph = body.paragraphs[0] if item_index == 0 else body.add_paragraph()
            paragraph.text = item
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(18)
            paragraph.font.color.rgb = INK
            paragraph.space_after = Pt(12)
        _footer(slide, len(prs.slides))


def _analysis_slides(prs: Presentation, analysis: dict[str, Any]) -> None:
    _flush_slide(prs, "数据范围与计算口径", [
        f"完整读取 {analysis['source_count']} 个文件、{analysis['sheet_count']} 个工作表，共 {analysis['row_count']:,} 行数据。",
        "数值由程序计算。空值与非数字排除，离群值仍计入统计；异常明细和全部数据保存在 Excel。",
        "环比使用上一个自然月，同比使用上一年同月。基期缺失或为 0 时不计算增长率。",
    ])
    for table in analysis["tables"]:
        for spec in table["charts"]:
            if not spec["labels"]:
                continue
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            labels, values = spec["labels"], spec["values"]
            suffix = ""
            if spec["kind"] == "line" and len(labels) > 24:
                labels, values = labels[-24:], values[-24:]
                suffix = "（最近 24 个月）"
            _heading(slide, spec["title"] + suffix)
            data = CategoryChartData()
            data.categories = labels
            data.add_series(spec["metric"], values)
            kind = XL_CHART_TYPE.LINE_MARKERS if spec["kind"] == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
            chart = slide.shapes.add_chart(kind, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.9), data).chart
            chart.has_legend = False
            chart.has_title = False
            chart.chart_style = 13 if spec["kind"] == "line" else 10
            chart.value_axis.tick_labels.font.size = Pt(11)
            chart.value_axis.minimum_scale = min(0, *values)
            if max(values) <= 0:
                chart.value_axis.maximum_scale = 0 if min(values) < 0 else 1
            if spec.get("value_format") == "percent":
                chart.value_axis.tick_labels.number_format = "0.0%"
            chart.category_axis.tick_labels.font.size = Pt(11)
            series = chart.series[0]
            if spec["kind"] == "line":
                series.format.line.color.rgb = TEAL
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = TEAL
            _footer(slide, len(prs.slides), f"来源：{table['source']} / {table['sheet']} · {table['row_count']} 行")


def markdown_to_pptx(md: str, dest: Path, analysis: dict[str, Any] | None = None) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover_title = "调研报告"
    sections: list[tuple[str, list[str]]] = []
    current_title: str | None = None
    current_bullets: list[str] = []
    saw_h1 = False

    def close_section() -> None:
        nonlocal current_title, current_bullets
        if current_title is not None:
            sections.append((current_title, list(current_bullets)))
        current_title = None
        current_bullets = []

    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        line = re.sub(r"`([^`]+)`", r"\1", line)
        if re.fullmatch(r"\s*\|[\s:|\-]+\|\s*", line):
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            line = " · ".join(cell.strip() for cell in line.strip().strip("|").split("|"))
        if line.startswith("# ") and not line.startswith("##"):
            text = line[2:].strip()
            if not saw_h1:
                cover_title = text or cover_title
                saw_h1 = True
            else:
                close_section()
                current_title = text or "章节"
            continue
        if line.startswith("## "):
            close_section()
            current_title = line[3:].strip() or "章节"
            continue
        if current_title is None:
            # body before any H2 → put under「摘要」
            current_title = "摘要"
        if line.startswith("### "):
            current_bullets.append(line[4:].strip())
        elif re.match(r"^[-*]\s+", line):
            current_bullets.append(re.sub(r"^[-*]\s+", "", line).strip())
        elif re.match(r"^\d+\.\s+", line):
            current_bullets.append(re.sub(r"^\d+\.\s+", "", line).strip())
        else:
            current_bullets.append(line.strip())

    close_section()

    # cover
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = cover_title
    cover.shapes.title.text_frame.paragraphs[0].font.color.rgb = INK
    cover.shapes.title.text_frame.paragraphs[0].font.size = Pt(34 if len(cover_title) < 45 else 24)
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = "AI办公搭子 · 自动生成"

    if not sections:
        _flush_slide(prs, "正文", [line.strip() for line in md.splitlines() if line.strip()] or ["（空报告）"])
    else:
        for title, bullets in sections:
            _flush_slide(prs, title, bullets)

    if analysis:
        _analysis_slides(prs, analysis)
    tmp = dest.with_name(f".{dest.stem}-{uuid.uuid4().hex}.pptx")
    try:
        prs.save(tmp)
        tmp.replace(dest)
    finally:
        tmp.unlink(missing_ok=True)
    return dest


def ensure_pptx_from_markdown(md_path: Path, pptx_path: Path) -> Path:
    md = md_path.read_text(encoding="utf-8")
    return markdown_to_pptx(md, pptx_path)
