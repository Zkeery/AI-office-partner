from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from app.core.errors import AppError
from app.services import table_analysis as tables
from app.services.pptx_export import markdown_to_pptx
from app.services.xlsx_export import analysis_to_xlsx, markdown_to_xlsx


def sales_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet.append(["日期", "地区", "销售额"] + [f"字段{i}" for i in range(4, 14)] + ["尾列"])
    for index in range(1, 101):
        sheet.append(["2025-09-01" if index <= 50 else "2026-09-01", "华东" if index % 2 else "华南", 777 if index == 100 else 100] + [index] * 11)
    second = book.create_sheet("退货")
    second.sheet_state = "hidden"
    second.append(["日期", "地区", "金额"])
    second.append(["2026-08-01", "华东", 20])
    second.append(["2026-08-02", "华南", 30])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def snapshot(tmp_path):
    source = tmp_path / "销售.xlsx"
    source.write_bytes(sales_workbook())
    identity = tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path / "artifacts", "汇总销售额")
    return tables.load_snapshot(tmp_path / "artifacts", identity)


def test_all_rows_columns_and_hidden_sheets_are_calculated(tmp_path):
    data = snapshot(tmp_path)
    summary = data["summary"]
    assert summary["complete"] is True
    assert summary["row_count"] == 102
    assert summary["sheet_count"] == 2
    main = summary["tables"][0]
    assert main["column_count"] == 14
    metrics = {metric["column"]: metric for metric in main["metrics"]}
    assert metrics["销售额"]["sum"] == 10677
    assert metrics["销售额"]["outlier_count"] == 1
    assert metrics["尾列"]["sum"] == 5050
    assert main["issues"][0]["row"] == 101
    assert summary["tables"][1]["metrics"][0]["sum"] == 50
    assert sum(main["groups"][0]["rows"][i]["value"] for i in range(2)) == 10677


def test_monthly_comparisons_and_invalid_values(tmp_path):
    source = tmp_path / "销售.csv"
    source.write_text("月份,地区,销售额\n2025-01,A,100\n2025-02,A,120\n2026-01,B,200\n2026-02,B,300\n2026-02,A,坏值\n2026-02,A,\n", encoding="utf-8")
    identity = tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path, "销售额")
    result = tables.load_snapshot(tmp_path, identity)["summary"]["tables"][0]
    metric = result["metrics"][0]
    assert metric["sum"] == 720
    assert metric["count"] == 4
    assert metric["invalid_count"] == metric["missing_count"] == 1
    periods = {row["period"]: row for row in result["periods"]}
    assert periods["2026-02"]["mom_pct"] == 50
    assert periods["2026-02"]["yoy_pct"] == 150
    assert periods["2026-01"]["mom_pct"] is None
    assert periods["2026-01"]["yoy_pct"] == 100
    assert result["issues"][0]["row"] == 6
    assert result["warnings"]


def test_zero_period_base_and_decimal_arithmetic(tmp_path):
    source = tmp_path / "销售.csv"
    source.write_text("月份,金额\n2025-02,0\n2026-01,0\n2026-02,0.1\n2026-02,0.2\n", encoding="utf-8")
    identity = tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path, "金额")
    result = tables.load_snapshot(tmp_path, identity)["summary"]["tables"][0]
    assert result["metrics"][0]["sum"] == 0.3
    assert result["periods"][-1]["mom_pct"] is None
    assert result["periods"][-1]["yoy_pct"] is None


def test_rates_are_averaged_in_groups_and_periods(tmp_path):
    source = tmp_path / "比例.csv"
    source.write_text("月份,地区,转化率\n2026-01,A,10%\n2026-01,A,30%\n2026-02,A,50%\n", encoding="utf-8")
    identity = tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path, "转化率")
    result = tables.load_snapshot(tmp_path, identity)["summary"]["tables"][0]
    assert result["metrics"][0]["aggregation"] == "mean"
    assert result["groups"][0]["rows"][0]["value"] == 0.3
    assert result["periods"][0]["value"] == 0.2
    assert result["periods"][1]["mom_pct"] == 150
    assert result["charts"][0]["value_format"] == "percent"


def test_uncalculated_formula_is_rejected_instead_of_becoming_zero(tmp_path):
    book = Workbook()
    book.active.append(["金额"])
    book.active.append(["=2+3"])
    source = tmp_path / "公式.xlsx"
    book.save(source)
    with pytest.raises(AppError) as failure:
        tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path, "汇总")
    assert failure.value.code == "TABLE_FORMULA_UNCALCULATED"
    assert "A2" in failure.value.message


def test_capacity_limit_is_explicit_not_sampling(tmp_path, monkeypatch):
    source = tmp_path / "长表.csv"
    source.write_text("金额\n1\n2\n3\n", encoding="utf-8")
    monkeypatch.setattr(tables, "MAX_ROWS", 2)
    with pytest.raises(AppError) as failure:
        tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path, "汇总")
    assert failure.value.code == "TABLE_TOO_LARGE"
    assert not list((tmp_path / "tables").glob("*.json"))


def test_excel_contains_full_data_formulas_and_native_charts(tmp_path):
    data = snapshot(tmp_path)
    output = analysis_to_xlsx(data, "# 销售结果\n\n正文", tmp_path / "result.xlsx")
    book = load_workbook(output)
    assert book["数据-销售"].max_row == 101
    assert book["数据-销售"].cell(101, 14).value == 100
    stats = list(book["指标汇总"].iter_rows(values_only=True))
    metric = next(row for row in stats if row[2] == "销售额")
    assert metric[6] == 10677
    assert metric[10] == "=SUM('数据-销售'!C2:C101)"
    assert len(book["分组统计"]._charts) == 2
    assert len(book["月度趋势"]._charts) == 2
    assert "数据-退货" in book.sheetnames
    book.close()


def test_user_text_is_not_written_as_an_excel_formula(tmp_path):
    source = tmp_path / "销售.csv"
    source.write_text('分类,金额\n"=HYPERLINK(""https://example.com"")",10\n普通,20\n', encoding="utf-8")
    identity = tables.build_snapshot([{"filename": source.name, "stored_path": str(source)}], tmp_path, "金额")
    data = tables.load_snapshot(tmp_path, identity)
    book_path = analysis_to_xlsx(data, "# 标题\n\n=1+1", tmp_path / "result.xlsx")
    book = load_workbook(book_path)
    assert book["数据-销售"]["A2"].data_type == "s"
    assert book["报告"]["A2"].data_type == "s"
    assert book["数据-销售"]["A2"].value.startswith("=HYPERLINK")
    book.close()


def test_ppt_uses_calculated_chart_data_and_does_not_drop_long_content(tmp_path):
    data = snapshot(tmp_path)
    markdown = "# 销售复盘\n\n## 说明\n\n" + "\n".join(f"- ITEM-{i:02d} " + "材料" * 130 + f" END-{i:02d}" for i in range(20))
    output = markdown_to_pptx(markdown, tmp_path / "result.pptx", analysis=data["summary"])
    presentation = Presentation(output)
    all_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
    assert "END-19" in all_text
    charts = [shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 4
    assert sum(charts[0].series[0].values) == 10677
    assert len(presentation.slides) > 5


def test_plain_report_markdown_table_exports_to_cells(tmp_path):
    output = markdown_to_xlsx("# 结果\n\n| 项目 | 数量 |\n| --- | --- |\n| A | 12 |\n| B | 34 |\n", tmp_path / "table.xlsx")
    book = load_workbook(output)
    assert book["表格1"]["B2"].value == 12
    assert book["表格1"]["B3"].value == 34
    book.close()
